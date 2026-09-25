import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

excel_path = r'C:\Codex_Personal\projects\HR\Китайский снабженец\简历汇总表_国际贸易专员.xlsx'
wb = openpyxl.load_workbook(excel_path)
ws = wb.active

headers = [str(cell.value).strip() if cell.value is not None else f"c{i}" for i, cell in enumerate(ws[2])]

target_ids = ['1', '8', '9', '12', '16']
for r in range(3, ws.max_row + 1):
    vals = [str(cell.value).strip() if cell.value is not None else '' for cell in ws[r]]
    d = dict(zip(headers, vals))
    cid = str(d.get('序号', ''))
    if cid in target_ids:
        print(f"ID {cid}: {d.get('姓名')} | {d.get('年龄/出生年月')} | {d.get('最高学历')}/{d.get('专业')} | {d.get('工作年限')} | {d.get('期望薪资')} | {d.get('期望城市')} | {d.get('联系电话')} {d.get('电子邮箱')}")
        print(f"   Role: {d.get('求职意向')}")
        print(f"   Exp: {d.get('工作经历摘要')}")
        print(f"   Note: {d.get('备注')}")
        print("-" * 70)
