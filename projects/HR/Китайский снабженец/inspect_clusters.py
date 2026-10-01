import sys
from pathlib import Path
import docx

sys.stdout.reconfigure(encoding="utf-8")

p = Path(r"D:\Документы Victus\Рабочее\Сотрудники\Дэвид\Equipment from China\For David abiut ckasters and competitors.docx")
out = Path(r"C:\Codex\projects\HR\Китайский снабженец\equipment_clusters_david.txt")

if p.exists():
    doc = docx.Document(str(p))
    lines = []
    for para in doc.paragraphs:
        if para.text.strip():
            lines.append(para.text.strip())
    for t in doc.tables:
        for r in t.rows:
            vals = [c.text.strip().replace('\n', ' ') for c in r.cells]
            lines.append(" | ".join(vals))
    txt = "\n\n".join(lines)
    out.write_text(txt, encoding="utf-8")
    print(f"Saved {len(lines)} lines to {out}")
    print("Content preview:")
    for l in lines[:20]:
        print(" ", l)
