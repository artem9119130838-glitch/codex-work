import sys
from pathlib import Path
import docx
from pypdf import PdfReader

# Ensure utf-8 output
sys.stdout.reconfigure(encoding="utf-8")

f1 = Path(r"D:\Документы Victus\Рабочее\Сотрудники\Дэвид\MSQIAN钱瑛个人简历-含俄文翻译.docx")
f2 = Path(r"D:\Документы Victus\Рабочее\Сотрудники\Дэвид\Jason Zuo 个人简历.pdf")

out_dir = Path(r"C:\Codex\projects\HR\Китайский снабженец")

print(f"File 1 exists: {f1.exists()} - {f1}")
print(f"File 2 exists: {f2.exists()} - {f2}")

# Parse File 1 (DOCX)
if f1.exists():
    doc = docx.Document(str(f1))
    lines = []
    for p in doc.paragraphs:
        if p.text.strip():
            lines.append(p.text.strip())
    for t in doc.tables:
        for r in t.rows:
            row_vals = [c.text.strip().replace('\n', ' ') for c in r.cells]
            # dedup contiguous duplicate cells (often merged cells in Word)
            dedup_row = []
            for v in row_vals:
                if not dedup_row or dedup_row[-1] != v:
                    dedup_row.append(v)
            if any(dedup_row):
                lines.append(" | ".join(dedup_row))
    text1 = "\n".join(lines)
    out1 = out_dir / "parsed_MSQIAN.txt"
    out1.write_text(text1, encoding="utf-8")
    print(f"Parsed MSQIAN: {len(lines)} lines written to {out1}")

# Parse File 2 (PDF)
if f2.exists():
    reader = PdfReader(str(f2))
    pdf_lines = []
    for idx, page in enumerate(reader.pages):
        page_text = page.extract_text()
        if page_text:
            pdf_lines.append(f"--- PAGE {idx+1} ---")
            pdf_lines.append(page_text.strip())
    text2 = "\n".join(pdf_lines)
    out2 = out_dir / "parsed_Jason_Zuo.txt"
    out2.write_text(text2, encoding="utf-8")
    print(f"Parsed Jason Zuo: {len(reader.pages)} pages, {len(text2)} chars written to {out2}")
