import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

excel_path = r'C:\Codex_Personal\projects\HR\Китайский снабженец\简历汇总表_国际贸易专员.xlsx'
wb = openpyxl.load_workbook(excel_path)
ws = wb.active

headers = [str(cell.value).strip() if cell.value is not None else f"c{i}" for i, cell in enumerate(ws[2])]

for r in range(3, ws.max_row + 1):
    vals = [str(cell.value).strip() if cell.value is not None else '' for cell in ws[r]]
    d = dict(zip(headers, vals))
    cid = d.get('序号', '')
    name = d.get('姓名', '')
    age = d.get('年龄/出生年月', '')
    edu = d.get('最高学历', '')
    major = d.get('专业', '')
    exp = d.get('工作年限', '')
    role = d.get('求职意向', '')
    salary = d.get('期望薪资', '')
    city = d.get('期望城市', '')
    history = d.get('工作经历摘要', '')
    notes = d.get('备注', '')
    contact = d.get('联系电话', '') or d.get('电子邮箱', '')
    print(f"ID {cid:2}: {name:6} | {age:5} | {edu:4}/{major:8} | {exp:10} | {salary:8} | {city:6} | {contact}")
    print(f"    Role: {role}")
    print(f"    Exp : {history[:140]}")
    print(f"    Note: {notes[:100]}")
    print("-" * 70)
