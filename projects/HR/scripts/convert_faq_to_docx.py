# -*- coding: utf-8 -*-
import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def convert_md_to_docx(md_path, docx_path):
    if not os.path.exists(md_path):
        print(f"Error: {md_path} not found")
        return False
        
    doc = docx.Document()
    
    # Page setup - 2cm margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
    # Styles
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()
        
    for line in lines:
        line_s = line.strip()
        if not line_s:
            continue
            
        if line_s.startswith("# "):
            # Title
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(line_s[2:].strip())
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
            
        elif line_s.startswith("## "):
            # Section Heading
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(line_s[3:].strip())
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x00, 0x40, 0x80)
            
        elif line_s.startswith("### "):
            # Question Heading
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line_s[4:].strip())
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
            
        elif line_s.startswith("> "):
            # Callout/quote
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line_s[2:].strip())
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            
        elif line_s.startswith("* **Поисковые синонимы / ключевые слова:**"):
            # Synonyms
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            r_label = p.add_run("Поисковые синонимы / ключевые слова: ")
            r_label.font.bold = True
            r_label.font.color.rgb = RGBColor(0x70, 0x40, 0x00)
            val = line_s.replace("* **Поисковые синонимы / ключевые слова:**", "").strip()
            r_val = p.add_run(val)
            r_val.font.color.rgb = RGBColor(0x70, 0x40, 0x00)
            r_val.font.italic = True
            
        elif line_s.startswith("* **Ответ:**") or line_s.startswith("**Ответ:**"):
            # Answer
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            r_label = p.add_run("Ответ: ")
            r_label.font.bold = True
            r_label.font.color.rgb = RGBColor(0x00, 0x66, 0x22)
            val = line_s.replace("* **Ответ:**", "").replace("**Ответ:**", "").strip()
            if val:
                p.add_run(val)
                
        elif line_s.startswith("- ") or line_s.startswith("* "):
            # Bullet list
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            txt = line_s[2:].strip()
            # Handle basic bold
            parts = txt.split("**")
            for i, pt in enumerate(parts):
                r = p.add_run(pt)
                if i % 2 == 1:
                    r.font.bold = True
                    
        elif line_s.startswith("---"):
            # Horizontal rule
            continue
            
        else:
            # Regular text
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            parts = line_s.split("**")
            for i, pt in enumerate(parts):
                r = p.add_run(pt)
                if i % 2 == 1:
                    r.font.bold = True

    doc.save(docx_path)
    print(f"Successfully generated DOCX: {docx_path}")
    return True

if __name__ == "__main__":
    md_file = os.path.join(os.getcwd(), "projects", "HR", "azat_supply_and_erp_faq.md")
    docx_file = os.path.join(os.getcwd(), "projects", "HR", "azat_supply_and_erp_faq.docx")
    convert_md_to_docx(md_file, docx_file)
