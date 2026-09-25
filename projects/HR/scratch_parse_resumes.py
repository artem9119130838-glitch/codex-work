import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

excel_path = r'C:\Codex_Personal\projects\HR\Китайский снабженец\简历汇总表_国际贸易专员.xlsx'
wb = openpyxl.load_workbook(excel_path)
ws = wb.active

# Row 2 contains actual column headers
headers = [str(cell.value).strip() if cell.value is not None else f"col_{c}" for c, cell in enumerate(ws[2], 1)]
print("Headers in Row 2:", headers)

candidates = []
for r in range(3, ws.max_row + 1):
    row_vals = [str(cell.value).strip() if cell.value is not None else '' for cell in ws[r]]
    if any(row_vals):
        cand_dict = {}
        for h, v in zip(headers, row_vals):
            cand_dict[h] = v
        candidates.append(cand_dict)

print(f"Total parsed candidates: {len(candidates)}")
for c in candidates:
    # Print key summary
    name = c.get('姓名', '')
    exp = c.get('工作年限', '')
    age = c.get('年龄', '')
    edu = c.get('学历', '')
    curr_job = c.get('目前职位', '') or c.get('期望职位', '')
    skills = c.get('专业技能', '') or c.get('核心优势', '') or c.get('工作经历摘要', '')
    print(f"ID {c.get('序号','')}: {name}, {age}岁, {edu}, {exp}经验. 职位: {curr_job}")
    for k, v in c.items():
        if k not in ['序号', '姓名', '年龄', '学历', '工作年限'] and v:
            print(f"   {k}: {v[:120]}")
    print("-" * 50)
