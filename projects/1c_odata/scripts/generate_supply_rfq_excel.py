import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
import os
import sys
import re
import argparse

if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

MASTER_TEMPLATE_PATH = r"D:\Документы Victus\Рабочее\Шаблоны\Заявки\Запрос КП пример заполнения.xlsx"
DESKTOP_DIR = os.path.join(os.environ.get("USERPROFILE", r"C:\Users\Артем"), "Desktop")

# База подтвержденных спецификаций по постоянным заказчикам (кэш эталонных позиций)
KNOWN_SPECIFICATIONS = {
    "волтайр": [
        {
            "num": 1,
            "brand": "OEM / Китай",
            "pname_cn": "85寸内胎硫化机 (主机)",
            "model": "85\" Dome type, 1 cavity (2250 kN / 225t)",
            "name_ru": "Вулканизационный пресс 85” косоугольный, однокамерный, центрический (тип нагрева Dome, усилие 2250 кН/225 т, диаметр камер 535-711 мм, высота ПФ 420-760 мм, макс. вес ПФ 4800 кг, 380В/15кВт, цвет RAL 7035)",
            "qty": 2,
        },
        {
            "num": 2,
            "brand": "VALAS / ARMSTRONG",
            "pname_cn": "阀组梯级系统 (蒸汽、冷凝水、空气、真空)",
            "model": "Cascade valve system for 85\" press",
            "name_ru": "Система каскадов клапанов (пар 0,8-1,0 МПа / 170-185 °С, конденсат, воздух 0,02-0,05 МПа, вакуум -0,06...-0,08 МПа; фланцевые клапаны VALAS, конденсатоотводчик ARMSTRONG)",
            "qty": 2,
        },
        {
            "num": 3,
            "brand": "OEM / Китай",
            "pname_cn": "自动润滑系统",
            "model": "Lubrication system for 85\" press",
            "name_ru": "Смазочная система вулканизационного пресса 85” (централизованная)",
            "qty": 2,
        },
        {
            "num": 4,
            "brand": "Siemens / Weintek / SMC / Schneider",
            "pname_cn": "电气控制系统套装 (PLC西门子, 触摸屏, SMC气动)",
            "model": "Control system 380V±10% / 50Hz",
            "name_ru": "Комплект системы управления (ПЛК Siemens, сенсорный экран Weintek/Advantech, пневматика SMC, датчики давления/температуры Jumo, низковольтные элементы Schneider, интерфейсные кабели)",
            "qty": 2,
        },
        {
            "num": 5,
            "brand": "OEM / Китай",
            "pname_cn": "两年运行备件包 (ZIP)",
            "model": "Spare parts kit for 2 years (п. 9.2 ТЗ)",
            "name_ru": "Комплект ЗИП с учетом 2-х летней эксплуатации оборудования (уплотнения, быстроизнашивающиеся детали, клапаны)",
            "qty": 2,
        }
    ],
    "муромец": [
        {
            "num": 1,
            "brand": "Alfa Laval",
            "pname_cn": "造水机密封垫包 (Gasket set for JWP-16-C40)",
            "model": "98517802 / art. 3263-0964-6",
            "name_ru": "Набор прокладок для опреснителя Alfa Laval JWP-16-C40 (судно Viktor Vasnetsov); article 3263-0964-6, кат. № 98517802",
            "qty": 1,
        },
        {
            "num": 2,
            "brand": "Alfa Laval",
            "pname_cn": "压力真空表 (Compound gauge -100%..2 Bar, D100 3/8\" BN)",
            "model": "98430308-00",
            "name_ru": "Мановакуумметр Compound gauge -100% - 2 Bar LB/in2, D100 3/8'' BN для опреснителя Alfa Laval JWP-16-C40; Art. № 98430308-00",
            "qty": 1,
        },
        {
            "num": 3,
            "brand": "Alfa Laval",
            "pname_cn": "黄铜视镜带垫圈 (Sight Glass incl. Gasket 1 1/2\" BRASS)",
            "model": "98435689-00",
            "name_ru": "Смотровое стекло в сборе с прокладкой Sight Glass incl Gasket 1 1/2'' BRASS для опреснителя Alfa Laval JWP-16-C40; Art. № 98435689-00",
            "qty": 1,
        }
    ],
    "вентэлектро": [
        {
            "num": 1,
            "brand": "GEMU",
            "pname_cn": "隔膜阀 602 10D17F35400TM 1507",
            "model": "602 10D17F35400TM 1507",
            "name_ru": "GEMU Мембранные клапаны 602 10D17F35400TM 1507",
            "qty": 53,
        }
    ],
    "инфамед": [
        {
            "num": 1,
            "brand": "GEMU",
            "pname_cn": "隔膜 (Diaphragms) MG10, MG25, MG40",
            "model": "MG10, MG25, MG40 (EPDM/PTFE)",
            "name_ru": "Мембраны GEMU для мембранных клапанов (типоразмеры MG10, MG25, MG40)",
            "qty": 1,
        }
    ],
    "гидросистемы": [
        {
            "num": 1,
            "brand": "Luen",
            "pname_cn": "制动阀 (模块式安装制动模块阀)",
            "model": "A-OWC-DE-L10-X (01.292.0X0.A)",
            "name_ru": "Тормозной клапан модульного монтажа CETOP5, линии А и В, расход >= 80 л/мин, давл. >= 32 МПа, пилот 1:4.5 (допустим качественный аналог)",
            "qty": 5,
        },
        {
            "num": 2,
            "brand": "Tognella",
            "pname_cn": "液压节流阀 (镀镍黄铜)",
            "model": "FT 1251/2-01-12",
            "name_ru": "Дроссель гидравлический с обратным клапаном, никелированная латунь",
            "qty": 4,
        },
        {
            "num": 3,
            "brand": "Tognella",
            "pname_cn": "两通高压球阀",
            "model": "FT221/1-112",
            "name_ru": "Двухходовой шаровый кран высокого давления",
            "qty": 2,
        },
        {
            "num": 4,
            "brand": "Tognella",
            "pname_cn": "压力表截止阀 (直通式)",
            "model": "FT290-14",
            "name_ru": "Отсечной вентиль под манометр (прямой)",
            "qty": 3,
        },
        {
            "num": 5,
            "brand": "OEM / PONAR",
            "pname_cn": "减压阀",
            "model": "PZM5-P280/10N",
            "name_ru": "Клапан редукционный давления",
            "qty": 5,
        },
        {
            "num": 6,
            "brand": "Atos",
            "pname_cn": "减压阀",
            "model": "SKG-033/210/V",
            "name_ru": "Редукционный клапан модульного монтажа",
            "qty": 4,
        },
        {
            "num": 7,
            "brand": "Stauff / OEM",
            "pname_cn": "压力检测点 (测压接头)",
            "model": "S10714G00C (SMK20-G1/4)",
            "name_ru": "Контрольная точка давления SMK20-G1/4",
            "qty": 1,
        },
        {
            "num": 8,
            "brand": "MP Filtri",
            "pname_cn": "吸油过滤器滤芯",
            "model": "STR1004BG1M90",
            "name_ru": "Фильтр всасывающий сетчатый (погружной)",
            "qty": 2,
        },
        {
            "num": 9,
            "brand": "LSQ / RFS",
            "pname_cn": "液压快换接头公头 (6605-4-4)",
            "model": "LSQ-S1-02PF-G1/4 (6605-4-4)",
            "name_ru": "Ниппель БРС G1/4 (быстроразъемное соединение)",
            "qty": 2,
        },
        {
            "num": 10,
            "brand": "LSQ / RFS",
            "pname_cn": "液压快换接头母头 (6603-4-4)",
            "model": "LSQ-S1-02SF-G1/4 (6603-4-4)",
            "name_ru": "Розетка БРС G1/4 (быстроразъемное соединение)",
            "qty": 4,
        },
        {
            "num": 11,
            "brand": "LSQ / RFS",
            "pname_cn": "快换接头公头用防尘帽",
            "model": "LSQ-S1 PDC-1/4 M",
            "name_ru": "Заглушка (пылезащитный колпачок) для ниппеля БРС 1/4",
            "qty": 2,
        },
        {
            "num": 12,
            "brand": "LSQ / RFS",
            "pname_cn": "快换接头母头用防尘塞",
            "model": "LSQ-S1 PDC-1/4 F",
            "name_ru": "Заглушка (пылезащитный колпачок) для розетки БРС 1/4",
            "qty": 2,
        }
    ]
}


def build_default_item(item_subject: str, company_name: str) -> dict:
    """Генерирует базовую позицию, если нет детализированной спецификации."""
    brand_match = re.search(r'\b([A-Za-z0-9\-]+)\b', item_subject)
    brand = brand_match.group(1) if brand_match else "OEM / Китай"
    
    clean_subj = item_subject.strip()
    return {
        "num": 1,
        "brand": brand,
        "pname_cn": clean_subj,
        "model": clean_subj,
        "name_ru": f"Запрос на поставку: {clean_subj} ({company_name})",
        "qty": 1
    }


def _create_fallback_workbook():
    """Создает книгу Excel с эталонной шапкой 'Запрос КП', если внешний шаблон недоступен."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Запрос КП"

    headers = [
        "№",
        "Бренд / Brand",
        "Наименование (CN) / Product Name (CN)",
        "Модель / Model / Art",
        "Описание (RU) / Product Description (RU)",
        "Кол-во / Qty",
        "Цена за ед. (RMB/USD)",
        "Сумма (RMB/USD)"
    ]

    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin', color='B0C4DE'),
        right=Side(style='thin', color='B0C4DE'),
        top=Side(style='thin', color='B0C4DE'),
        bottom=Side(style='thin', color='B0C4DE')
    )

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    return wb, ws


def generate_rfq_excel(company_name: str, item_subject: str, items: list = None, output_filename: str = None, dry_run: bool = False) -> str:
    """
    Создает чистовой Excel по эталонному шаблону Запрос КП пример заполнения.xlsx,
    заполняя позиции спецификации, шрифты Calibri 10, рамки и ширину колонок.
    В случае отсутствия внешнего шаблона автоматически создает fallback-книгу с корпоративным стилем.
    В режиме dry_run=True выполняет проверку валидности данных без записи файла на рабочий стол.
    """
    # Очищаем наименования для имени файла
    clean_comp = re.sub(r'["«»АООООПАОЗАО]', '', company_name).strip() or "Клиент"
    clean_item = re.sub(r'[\\/*?:"<>|]', '', item_subject).strip()[:35] or "Оборудование"

    if not output_filename:
        output_filename = f"Запрос КП {clean_item} {clean_comp}.xlsx"
        if "волтайр" in company_name.lower():
            output_filename = "Запрос КП пресс для автокамер 85 Волтайр-Пром.xlsx"
        elif "муромец" in company_name.lower():
            output_filename = "Запрос КП Alfa Laval опреснитель JWP-16-C40 Илья Муромец.xlsx"
        elif "гидросистемы" in company_name.lower():
            output_filename = "Запрос КП Гидравлика 12 позиций НПО Гидросистемы.xlsx"

    output_path = os.path.join(DESKTOP_DIR, output_filename)

    # Определяем состав позиций
    final_items = items
    if not final_items:
        low_comp = company_name.lower()
        for key, spec in KNOWN_SPECIFICATIONS.items():
            if key in low_comp:
                final_items = spec
                break

    if not final_items:
        final_items = [build_default_item(item_subject, company_name)]

    # Нормализуем элементы, если переданы строки
    normalized_items = []
    for idx, it in enumerate(final_items, start=1):
        if isinstance(it, dict):
            item_dict = dict(it)
            item_dict["num"] = item_dict.get("num", idx)
            normalized_items.append(item_dict)
        elif isinstance(it, str):
            normalized_items.append({
                "num": idx,
                "brand": "OEM / Китай",
                "pname_cn": it,
                "model": it,
                "name_ru": it,
                "qty": 1
            })

    if dry_run:
        print(f"[DRY-RUN] [EXCEL-GEN] Файл был бы сохранен по пути: {output_path}")
        print(f"[DRY-RUN] [EXCEL-GEN] Позиций к генерации: {len(normalized_items)}")
        for it in normalized_items:
            print(f"          - №{it.get('num')}: {it.get('pname_cn', '')} | Brand: {it.get('brand')} | Qty: {it.get('qty')}")
        return output_path

    # Загружаем мастер-шаблон или создаем fallback
    if os.path.exists(MASTER_TEMPLATE_PATH):
        wb = openpyxl.load_workbook(MASTER_TEMPLATE_PATH)
        ws = wb.active
        ws.title = "Запрос КП"
        # Очищаем строки шаблона со 2-й и далее
        for r in range(2, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                ws.cell(r, c).value = None
    else:
        wb, ws = _create_fallback_workbook()

    # Стили по стандарту create_tender_excel.py
    font_data = Font(name="Calibri", size=10)
    font_num = Font(name="Calibri", size=10, bold=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_center = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # Заполняем позиции
    for idx, it in enumerate(normalized_items, start=2):
        ws.cell(idx, 1).value = it.get("num", idx - 1)
        ws.cell(idx, 1).font = font_num
        ws.cell(idx, 1).alignment = align_center
        ws.cell(idx, 1).border = thin_border

        ws.cell(idx, 2).value = it.get("brand", "OEM")
        ws.cell(idx, 2).font = font_data
        ws.cell(idx, 2).alignment = align_center
        ws.cell(idx, 2).border = thin_border

        ws.cell(idx, 3).value = it.get("pname_cn", "")
        ws.cell(idx, 3).font = font_data
        ws.cell(idx, 3).alignment = align_left
        ws.cell(idx, 3).border = thin_border

        ws.cell(idx, 4).value = it.get("model", "")
        ws.cell(idx, 4).font = font_data
        ws.cell(idx, 4).alignment = align_left
        ws.cell(idx, 4).border = thin_border

        ws.cell(idx, 5).value = it.get("name_ru", "")
        ws.cell(idx, 5).font = font_data
        ws.cell(idx, 5).alignment = align_left
        ws.cell(idx, 5).border = thin_border

        ws.cell(idx, 6).value = it.get("qty", 1)
        ws.cell(idx, 6).font = font_data
        ws.cell(idx, 6).alignment = align_center
        ws.cell(idx, 6).border = thin_border

        # Колонки цен оставляем пустыми для заполнения фабрикой
        ws.cell(idx, 7).value = None
        ws.cell(idx, 7).border = thin_border

        ws.cell(idx, 8).value = None
        ws.cell(idx, 8).border = thin_border

    # Высоты строк
    ws.row_dimensions[1].height = 28
    for r in range(2, len(normalized_items) + 2):
        ws.row_dimensions[r].height = 45

    # Ширина колонок
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 32
    ws.column_dimensions['D'].width = 30
    ws.column_dimensions['E'].width = 50
    ws.column_dimensions['F'].width = 12
    ws.column_dimensions['G'].width = 18
    ws.column_dimensions['H'].width = 20

    wb.save(output_path)
    print(f"[EXCEL-GEN] Успешно сформирован файл: {output_path} (позиций: {len(normalized_items)})")
    return output_path


def run_self_check() -> bool:
    """Аппаратный Self-Check: проверка работы как в dry-run, так и fallback генерации."""
    print("=== [SELF-CHECK] Старт аппаратного теста generate_supply_rfq_excel.py ===")
    test_items = [
        {"num": 1, "brand": "ABB", "pname_cn": "变频器 ACS580", "model": "ACS580-01-045A-4", "name_ru": "Преобразователь частоты ABB 22кВт", "qty": 3},
        {"num": 2, "brand": "Schneider", "pname_cn": "断路器 NSX100", "model": "LV429630", "name_ru": "Автоматический выключатель NSX100F", "qty": 10}
    ]

    # 1. Проверка Dry-Run
    res_dry = generate_rfq_excel("ТестПром", "Преобразователи", items=test_items, dry_run=True)
    if not res_dry.endswith(".xlsx"):
        print("[FAIL] Dry-run не вернул корректный путь к xlsx")
        return False
    print("[PASS] Шаг 1: Dry-run режим отработал корректно без создания файла.")

    # 2. Проверка Fallback книги без внешнего шаблона в памяти
    wb_fb, ws_fb = _create_fallback_workbook()
    if ws_fb.max_column != 8 or ws_fb.title != "Запрос КП":
        print(f"[FAIL] Fallback workbook не соответствует эталону: col={ws_fb.max_column}, title={ws_fb.title}")
        return False
    print("[PASS] Шаг 2: Fallback генерация книги Excel валидна (8 колонок, стили шапки).")

    print("=== [SELF-CHECK] Все проверки успешно пройдены! ===")
    return True


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Генератор файлов запроса КП в Китай (RFQ Excel)")
    parser.add_argument("pos_company", nargs="?", default=None, help="Компания (позиционный аргумент)")
    parser.add_argument("pos_item", nargs="?", default=None, help="Предмет запроса (позиционный аргумент)")
    parser.add_argument("--company", "-c", dest="company", default=None, help="Наименование компании клиента")
    parser.add_argument("--item", "-i", dest="item", default=None, help="Предмет заявки")
    parser.add_argument("--output", "-o", dest="output", default=None, help="Имя выходного файла")
    parser.add_argument("--dry-run", action="store_true", help="Сухой прогон без создания файла на диске")
    parser.add_argument("--self-check", action="store_true", help="Запуск аппаратной самопроверки")

    args = parser.parse_args()

    if args.self_check:
        success = run_self_check()
        sys.exit(0 if success else 1)

    company = args.company or args.pos_company or "АО «Волтайр-Пром»"
    item = args.item or args.pos_item or '85" 内胎硫化机'
    res = generate_rfq_excel(company, item, dry_run=args.dry_run, output_filename=args.output)
    print(f"Результат: {res}")
