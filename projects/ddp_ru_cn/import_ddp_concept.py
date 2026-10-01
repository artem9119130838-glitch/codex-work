import sys
from pathlib import Path
import docx

sys.stdout.reconfigure(encoding="utf-8")

src = Path(r"D:\Документы Victus\Рабочее\В работе\Проект DDP Ru-CN\концепция DDP сотрудничества (рус).docx")
dest_dir = Path(r"C:\Codex\projects\ddp_ru_cn")
dest_dir.mkdir(parents=True, exist_ok=True)

if src.exists():
    doc = docx.Document(str(src))
    lines = []
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            lines.append(txt)
    for t in doc.tables:
        for r in t.rows:
            vals = [c.text.strip().replace('\n', ' ') for c in r.cells]
            dedup = []
            for v in vals:
                if not dedup or dedup[-1] != v:
                    dedup.append(v)
            if any(dedup):
                lines.append(" | ".join(dedup))
    
    out_md = dest_dir / "CONCEPT_DDP_COLLABORATION_RU.md"
    out_md.write_text("\n\n".join(lines), encoding="utf-8")
    print(f"Saved DDP concept to {out_md} ({len(lines)} lines)")
else:
    print(f"Source file not found: {src}")
