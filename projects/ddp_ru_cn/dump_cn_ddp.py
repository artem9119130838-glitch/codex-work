import sys
from pathlib import Path
import docx

sys.stdout.reconfigure(encoding="utf-8")

p = Path(r"D:\Документы Victus\Рабочее\Сотрудники\Дэвид\Проект DDP Ru-CN\DDP Russia for Chinese Partners.docx")
doc = docx.Document(str(p))

lines = []
for idx, para in enumerate(doc.paragraphs):
    t = para.text.strip()
    if t:
        lines.append(f"[{idx}] {t}")

for t in doc.tables:
    for r in t.rows:
        row_txt = " | ".join([c.text.strip().replace('\n', ' ') for c in r.cells])
        lines.append(f"[TABLE] {row_txt}")

out = Path(r"C:\Codex\projects\ddp_ru_cn\DDP_Russia_for_Chinese_Partners_dump.txt")
out.write_text("\n".join(lines), encoding="utf-8")
print(f"Dumped {len(lines)} lines to {out}")
