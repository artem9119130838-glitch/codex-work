import os
import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

def read_doc_text(path):
    if not os.path.exists(path):
        return ""
    d = docx.Document(path)
    lines = []
    for p in d.paragraphs:
        t = p.text.strip()
        if t:
            lines.append(t)
    for table in d.tables:
        for row in table.rows:
            row_txt = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
            if any(row_txt):
                lines.append(" | ".join(row_txt))
    return "\n".join(lines)

notes_dir = r"D:\Документы Victus\Рабочее\В работе\SEO GPT Проект\Важные заметки"
for f in os.listdir(notes_dir):
    if f.lower().startswith("архив"): continue
    fp = os.path.join(notes_dir, f)
    if f.endswith(".docx"):
        txt = read_doc_text(fp)
        print(f"\n==================== {f} ({len(txt)} chars) ====================")
        print(txt[:700] + ("..." if len(txt) > 700 else ""))
    elif f.endswith(".csv"):
        print(f"\n==================== {f} (CSV file, {os.path.getsize(fp)} bytes) ====================")
