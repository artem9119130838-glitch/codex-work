import os
import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

def extract_key_sections(path):
    d = docx.Document(path)
    print("=== " + os.path.basename(path) + " ===")
    for i, p in enumerate(d.paragraphs):
        t = p.text.strip()
        if any(w in t.lower() for w in ['главный итог', 'приоритеты', 'цель проекта', 'p0', 'p1', 'архитектура']):
            if len(t) < 120:
                print(f"\n--- Section: {t} ---")
                # print next 3 paras
                for j in range(1, 4):
                    if i + j < len(d.paragraphs):
                        next_t = d.paragraphs[i+j].text.strip()
                        if next_t:
                            print(f"  {next_t[:150]}...")

if __name__ == '__main__':
    extract_key_sections(r"D:\Документы Victus\Рабочее\В работе\SEO GPT Проект\Итоговый отчёт для SEO-аудита CHATGPT.docx")
    extract_key_sections(r"D:\Документы Victus\Рабочее\В работе\SEO GPT Проект\Gemini Резюме проекта Контентная фабрика Longwang.ru (Кейсы и Экспертные статьи).docx")
