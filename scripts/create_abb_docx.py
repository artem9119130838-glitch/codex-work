import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os

doc = docx.Document()

# Page setup: Margins
for section in doc.sections:
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

# Color Palette: Corporate Navy / Slate
COLOR_NAVY = RGBColor(0, 51, 102)
COLOR_TEXT = RGBColor(33, 37, 41)
COLOR_MUTED = RGBColor(108, 117, 125)
COLOR_BG_HEADER = '003366'
COLOR_BG_ROW_ALT = 'F8F9FA'

def style_paragraph(p, space_before=0, space_after=6, line_spacing=1.15):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    style_paragraph(p, space_before=14, space_after=6)
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = COLOR_NAVY
    return p

def set_cell_shading(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/><w:left w:val="none"/><w:right w:val="none"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/><w:insideV w:val="none"/></w:tblBorders>')
    tblPr.append(borders)

# Title Block
p_title = doc.add_paragraph()
style_paragraph(p_title, space_before=0, space_after=3)
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_title = p_title.add_run('ТЕХНИЧЕСКИЙ ПАСПОРТ И ОПИСАНИЕ ОБОРУДОВАНИЯ')
r_title.font.name = 'Arial'
r_title.font.size = Pt(15)
r_title.font.bold = True
r_title.font.color.rgb = COLOR_NAVY

p_sub = doc.add_paragraph()
style_paragraph(p_sub, space_before=0, space_after=12)
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_sub = p_sub.add_run('Для целей таможенного оформления и идентификации по ТН ВЭД ЕАЭС\nПоставка оборудования концерна ABB (Швейцария / Италия / Финляндия)')
r_sub.font.name = 'Arial'
r_sub.font.size = Pt(10)
r_sub.font.color.rgb = COLOR_MUTED

# General Info Box
p_intro = doc.add_paragraph()
style_paragraph(p_intro, space_before=0, space_after=8)
r_intro = p_intro.add_run(
    'Настоящий документ содержит подробные технические характеристики, электрофизические параметры, '
    'габаритные размеры и описание назначения силовой коммутационной аппаратуры низкого напряжения (до 1000 В) '
    'производства концерна ABB. Оборудование предназначено для установки в промышленных шкафах распределения (ГРЩ, ВРУ) '
    'и защиты силовых цепей переменного тока от перегрузок и токов короткого замыкания.'
)
r_intro.font.name = 'Arial'
r_intro.font.size = Pt(9.5)

# Special Notice Box
p_note = doc.add_paragraph()
style_paragraph(p_note, space_before=2, space_after=12)
r_note_b = p_note.add_run('Официальное подтверждение для таможенного контроля: ')
r_note_b.font.name = 'Arial'
r_note_b.font.size = Pt(9)
r_note_b.font.bold = True
r_note_b.font.color.rgb = COLOR_NAVY
r_note_t = p_note.add_run(
    'Все указанные позиции представляют собой чисто электротехнические аппараты. Оборудование НЕ содержит в своем составе '
    'радиоэлектронных средств (РЭС), высокочастотных устройств (ВЧУ), беспроводных модулей передачи данных (Wi-Fi, Bluetooth, GSM) '
    'и криптографических (шифровальных) средств. Ввоз не требует оформления нотификации ФСБ и лицензий Минпромторга РФ. '
    'Оборудование не является медицинским изделием (РУ Минздрава не требуется).'
)
r_note_t.font.name = 'Arial'
r_note_t.font.size = Pt(9)

# ----------------- ITEM 1 -----------------
add_heading_1(doc, '1. Автоматический выключатель в литом корпусе ABB XT3S 250')
p_i1 = doc.add_paragraph()
style_paragraph(p_i1, space_before=2, space_after=6)
r_i1 = p_i1.add_run('Полное наименование / артикул: XT3S 250 MA100-600-1200 PMP 3P\nКод ТН ВЭД ЕАЭС: 8536 20 900 0 (Автоматические выключатели на напряжение до 1000 В, силу тока более 63 А)')
r_i1.font.name = 'Arial'
r_i1.font.size = Pt(9.5)
r_i1.font.bold = True

t1 = doc.add_table(rows=1, cols=2)
t1.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t1, color='FFFFFF')
cell_img1, cell_desc1 = t1.rows[0].cells
cell_img1.width = Inches(2.0)
cell_desc1.width = Inches(4.8)

p_img1 = cell_img1.paragraphs[0]
style_paragraph(p_img1, space_before=0, space_after=0)
p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
if os.path.exists(r'c:\Codex\temp_assets\xt3s.png'):
    p_img1.add_run().add_picture(r'c:\Codex\temp_assets\xt3s.png', width=Inches(1.6))

p_d1 = cell_desc1.paragraphs[0]
style_paragraph(p_d1, space_before=0, space_after=4)
p_d1.add_run('Назначение и применение:\nАвтоматический выключатель в литом корпусе серии SACE Tmax XT3. '
             'Предназначен для защиты электрических двигателей и силовых распределительных сетей от токов короткого замыкания. '
             'Оснащен электромагнитным расцепителем типа MA с регулируемым порогом срабатывания от КЗ (600...1200 А). '
             'Исполнение PMP — стационарное с передними выводами, 3-полюсное.').font.size = Pt(8.5)

t1_specs = doc.add_table(rows=8, cols=2)
t1_specs.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t1_specs, color='DDDDDD')
data_1 = [
    ('Параметр', 'Значение'),
    ('Номинальный непрерывный ток (In / Iu)', '100 А (типоразмер каркаса до 250 А)'),
    ('Номинальное рабочее напряжение (Ue)', '690 В перем. тока (AC) / 500 В пост. тока (DC)'),
    ('Номинальное напряжение изоляции (Ui)', '800 В'),
    ('Предельная отключающая способность (Icu при 415 В)', '50 кА (класс S - повышенная стойкость)'),
    ('Тип расцепителя', 'Электромагнитный (MA 100, уставка 600...1200 А)'),
    ('Число полюсов / исполнение', '3 полюса (3P) / стационарное с передними выводами'),
    ('Габариты (Ш х В х Г) / Вес 1 шт', '105 х 150 х 70 мм / 1,65 кг (нетто), ~1,9 кг (брутто)')
]
for idx, (param, val) in enumerate(data_1):
    row = t1_specs.rows[idx]
    c0, c1 = row.cells
    c0.width = Inches(2.6)
    c1.width = Inches(4.2)
    set_cell_margins(c0, 50, 50, 90, 90)
    set_cell_margins(c1, 50, 50, 90, 90)
    if idx == 0:
        set_cell_shading(c0, COLOR_BG_HEADER)
        set_cell_shading(c1, COLOR_BG_HEADER)
        r0 = c0.paragraphs[0].add_run(param)
        r1 = c1.paragraphs[0].add_run(val)
        r0.font.bold = r1.font.bold = True
        r0.font.color.rgb = r1.font.color.rgb = RGBColor(255, 255, 255)
    else:
        if idx % 2 == 1:
            set_cell_shading(c0, COLOR_BG_ROW_ALT)
            set_cell_shading(c1, COLOR_BG_ROW_ALT)
        r0 = c0.paragraphs[0].add_run(param)
        r1 = c1.paragraphs[0].add_run(val)
        r0.font.bold = True
    c0.paragraphs[0].runs[0].font.size = Pt(8.5)
    c1.paragraphs[0].runs[0].font.size = Pt(8.5)

# ----------------- ITEM 2 -----------------
add_heading_1(doc, '2. Автоматический выключатель в литом корпусе ABB T6N 630')
p_i2 = doc.add_paragraph()
style_paragraph(p_i2, space_before=2, space_after=6)
r_i2 = p_i2.add_run('Полное наименование / артикул: T6N-630 PR221DS-LSI R630 FF 3P\nКод ТН ВЭД ЕАЭС: 8536 20 900 0 (Автоматические выключатели на напряжение до 1000 В, силу тока более 63 А)')
r_i2.font.name = 'Arial'
r_i2.font.size = Pt(9.5)
r_i2.font.bold = True

t2 = doc.add_table(rows=1, cols=2)
t2.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t2, color='FFFFFF')
cell_img2, cell_desc2 = t2.rows[0].cells
cell_img2.width = Inches(2.0)
cell_desc2.width = Inches(4.8)

p_img2 = cell_img2.paragraphs[0]
style_paragraph(p_img2, space_before=0, space_after=0)
p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
if os.path.exists(r'c:\Codex\temp_assets\t6n.jpg'):
    p_img2.add_run().add_picture(r'c:\Codex\temp_assets\t6n.jpg', width=Inches(1.6))

p_d2 = cell_desc2.paragraphs[0]
style_paragraph(p_d2, space_before=0, space_after=4)
p_d2.add_run('Назначение и применение:\nМощный силовой автоматический выключатель серии SACE Tmax T6. '
             'Предназначен для защиты вводных и распределительных линий главных распределительных щитов (ГРЩ). '
             'Оснащен электронным расцепителем PR221DS-LSI с селективной защитой от перегрузки (L), селективной задержкой короткого замыкания (S) '
             'и мгновенной токовой отсечкой (I). Настройка параметров производится DIP-переключателями на лицевой панели.').font.size = Pt(8.5)

t2_specs = doc.add_table(rows=8, cols=2)
t2_specs.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t2_specs, color='DDDDDD')
data_2 = [
    ('Параметр', 'Значение'),
    ('Номинальный непрерывный ток (In / Iu)', '630 А'),
    ('Номинальное рабочее напряжение (Ue)', '690 В перем. тока (AC)'),
    ('Номинальное напряжение изоляции (Ui)', '1000 В'),
    ('Предельная отключающая способность (Icu при 415 В)', '36 кА (класс N - стандартная отключающая способность)'),
    ('Тип расцепителя', 'Электронный микропроцессорный PR221DS-LSI (без радиочастотных и криптомодулей)'),
    ('Число полюсов / исполнение', '3 полюса (3P) / стационарное с передними выводами (FF)'),
    ('Габариты (Ш х В х Г) / Вес 1 шт', '210 х 268 х 103,5 мм / 9,80 кг (нетто), ~11,0 кг (брутто)')
]
for idx, (param, val) in enumerate(data_2):
    row = t2_specs.rows[idx]
    c0, c1 = row.cells
    c0.width = Inches(2.6)
    c1.width = Inches(4.2)
    set_cell_margins(c0, 50, 50, 90, 90)
    set_cell_margins(c1, 50, 50, 90, 90)
    if idx == 0:
        set_cell_shading(c0, COLOR_BG_HEADER)
        set_cell_shading(c1, COLOR_BG_HEADER)
        r0 = c0.paragraphs[0].add_run(param)
        r1 = c1.paragraphs[0].add_run(val)
        r0.font.bold = r1.font.bold = True
        r0.font.color.rgb = r1.font.color.rgb = RGBColor(255, 255, 255)
    else:
        if idx % 2 == 1:
            set_cell_shading(c0, COLOR_BG_ROW_ALT)
            set_cell_shading(c1, COLOR_BG_ROW_ALT)
        r0 = c0.paragraphs[0].add_run(param)
        r1 = c1.paragraphs[0].add_run(val)
        r0.font.bold = True
    c0.paragraphs[0].runs[0].font.size = Pt(8.5)
    c1.paragraphs[0].runs[0].font.size = Pt(8.5)

# ----------------- ITEM 3 -----------------
add_heading_1(doc, '3. Выключатель нагрузки (рубильник) ABB OT800E03')
p_i3 = doc.add_paragraph()
style_paragraph(p_i3, space_before=2, space_after=6)
r_i3 = p_i3.add_run('Полное наименование / артикул: OT800E03 (Заказной код: 1SCA022753R5920)\nКод ТН ВЭД ЕАЭС: 8536 50 800 0 (Переключатели прочие на напряжение не более 1000 В)')
r_i3.font.name = 'Arial'
r_i3.font.size = Pt(9.5)
r_i3.font.bold = True

t3 = doc.add_table(rows=1, cols=2)
t3.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t3, color='FFFFFF')
cell_img3, cell_desc3 = t3.rows[0].cells
cell_img3.width = Inches(2.0)
cell_desc3.width = Inches(4.8)

p_img3 = cell_img3.paragraphs[0]
style_paragraph(p_img3, space_before=0, space_after=0)
p_img3.alignment = WD_ALIGN_PARAGRAPH.CENTER
if os.path.exists(r'c:\Codex\temp_assets\ot800.jpg'):
    p_img3.add_run().add_picture(r'c:\Codex\temp_assets\ot800.jpg', width=Inches(1.6))

p_d3 = cell_desc3.paragraphs[0]
style_paragraph(p_d3, space_before=0, space_after=4)
p_d3.add_run('Назначение и применение:\nТрехполюсный силовой выключатель нагрузки (рубильник-разъединитель) с ручным управлением серии OT800. '
             'Предназначен для включения, пропускания и отключения номинальных токов под нагрузкой, а также создания видимого безопасного разрыва '
             'цепи при техническом обслуживании электрооборудования. Монтируется на монтажную панель в распределительных шкафах.').font.size = Pt(8.5)

t3_specs = doc.add_table(rows=8, cols=2)
t3_specs.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t3_specs, color='DDDDDD')
data_3 = [
    ('Параметр', 'Значение'),
    ('Номинальный тепловой ток на открытом воздухе (Ith)', '800 А (при 40 °C)'),
    ('Номинальный рабочий ток (Ie при AC-21A / AC-22A)', '800 А (до 690 В)'),
    ('Номинальное рабочее напряжение (Ue)', '1000 В перем. тока (AC)'),
    ('Номинальное напряжение изоляции (Ui)', '1000 В'),
    ('Номинальный кратковременно выдерживаемый ток (Icw)', '20 кА в течение 1 сек'),
    ('Число полюсов / тип механизма', '3 полюса (3P) / ручной привод (базовый блок)'),
    ('Габариты (Ш х В х Г) / Вес 1 шт', '266 х 250 х 140 мм / 6,30 кг (нетто), ~7,0 кг (брутто)')
]
for idx, (param, val) in enumerate(data_3):
    row = t3_specs.rows[idx]
    c0, c1 = row.cells
    c0.width = Inches(2.6)
    c1.width = Inches(4.2)
    set_cell_margins(c0, 50, 50, 90, 90)
    set_cell_margins(c1, 50, 50, 90, 90)
    if idx == 0:
        set_cell_shading(c0, COLOR_BG_HEADER)
        set_cell_shading(c1, COLOR_BG_HEADER)
        r0 = c0.paragraphs[0].add_run(param)
        r1 = c1.paragraphs[0].add_run(val)
        r0.font.bold = r1.font.bold = True
        r0.font.color.rgb = r1.font.color.rgb = RGBColor(255, 255, 255)
    else:
        if idx % 2 == 1:
            set_cell_shading(c0, COLOR_BG_ROW_ALT)
            set_cell_shading(c1, COLOR_BG_ROW_ALT)
        r0 = c0.paragraphs[0].add_run(param)
        r1 = c1.paragraphs[0].add_run(val)
        r0.font.bold = True
    c0.paragraphs[0].runs[0].font.size = Pt(8.5)
    c1.paragraphs[0].runs[0].font.size = Pt(8.5)

# Save to paths
out_victus = r'D:\Документы Victus\Рабочее\КП\Готовые КП\Техническое_описание_оборудования_ABB.docx'
out_local = r'c:\Codex\Техническое_описание_оборудования_ABB.docx'

os.makedirs(os.path.dirname(out_victus), exist_ok=True)
doc.save(out_victus)
doc.save(out_local)
print('Successfully saved to:', out_victus, 'and', out_local)
